"""多模态转文字：图片 → OCR+描述；音频 → 分段转写；视频 → 音轨转写 + 关键帧描述。

qwen3-asr-flash 单次最长 5 分钟、10MB，所以长录音先用 ffmpeg 切成 4 分钟一段
（16kHz 单声道 32kbps，约 1MB/段），并行转写后按顺序拼接。
"""

import base64
import json
import mimetypes
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import httpx

SEGMENT_SECONDS = 240
MAX_FRAMES = 8

IMAGE_PROMPT = (
    "这是用户调研中受访者提供的图片（可能是截图、照片、表格、聊天记录）。\n"
    "1. 先完整提取图中所有文字（保持原有顺序/表格结构）；\n"
    "2. 再用 2~4 句话描述图片内容，以及它反映了受访者怎样的工作/生活场景。\n"
    "只描述看到的，不要推测图中没有的信息。"
)
FRAMES_PROMPT = (
    "以下是用户调研视频的关键帧（按时间顺序，标注了时间点）。\n"
    "逐帧简述画面：受访者在做什么、用的什么工具/软件、屏幕上的关键文字。\n"
    "最后用 1~2 句话总结视频展示的操作流程。只描述看到的。"
)


class MediaProcessor:
    def __init__(self, api_key: str, base_url: str, asr_model: str, vl_model: str,
                 http: httpx.Client | None = None, workers: int = 4):
        self.asr_model = asr_model
        self.vl_model = vl_model
        self.workers = workers
        self._http = http or httpx.Client(
            base_url=base_url, timeout=300, trust_env=False,
            headers={"Authorization": f"Bearer {api_key}"})

    # ---------- 入口 ----------
    def extract(self, filename: str, data: bytes, mime: str | None) -> dict:
        kind = media_kind(filename, mime)
        try:
            with tempfile.TemporaryDirectory() as tmp:
                src = Path(tmp) / ("in" + (Path(filename).suffix or ".bin"))
                src.write_bytes(data)
                if kind == "image":
                    text = self.describe_images([(data, mime or "image/png", None)], IMAGE_PROMPT)
                elif kind == "audio":
                    text = self.transcribe(src, Path(tmp))
                elif kind == "video":
                    text = self._video(src, Path(tmp))
                elif kind == "document":
                    text = read_document(src)
                else:
                    return {"filename": filename, "kind": kind, "error": f"不支持的类型: {mime}"}
            return {"filename": filename, "kind": kind, "text": text}
        except Exception as e:  # noqa: BLE001  单个文件失败不影响其它文件
            return {"filename": filename, "kind": kind, "error": f"{type(e).__name__}: {e}"}

    def _video(self, src: Path, tmp: Path) -> str:
        parts = []
        if has_audio(src):
            parts.append("【音轨转写】\n" + (self.transcribe(src, tmp) or "（无语音内容）"))
        frames = extract_frames(src, tmp)
        if frames:
            imgs = [(f.read_bytes(), "image/jpeg", ts) for f, ts in frames]
            parts.append("【画面描述】\n" + self.describe_images(imgs, FRAMES_PROMPT))
        return "\n\n".join(parts)

    # ---------- 音频 ----------
    def transcribe(self, src: Path, tmp: Path) -> str:
        segs = split_audio(src, tmp)
        with ThreadPoolExecutor(self.workers) as ex:
            texts = list(ex.map(self._asr_one, segs))
        lines = []
        for i, t in enumerate(texts):
            if t.strip():
                lines.append(f"[{fmt_ts(i * SEGMENT_SECONDS)}] {t.strip()}")
        return "\n".join(lines)

    def _asr_one(self, seg: Path) -> str:
        uri = "data:audio/mpeg;base64," + base64.b64encode(seg.read_bytes()).decode()
        r = self._http.post("/chat/completions", json={
            "model": self.asr_model,
            "messages": [{"role": "user", "content": [
                {"type": "input_audio", "input_audio": {"data": uri}}]}],
            "stream": False,
        })
        r.raise_for_status()
        return r.json()["choices"][0]["message"]["content"] or ""

    # ---------- 视觉 ----------
    def describe_images(self, images: list[tuple[bytes, str, float | None]], prompt: str) -> str:
        content: list[dict] = []
        for data, mime, ts in images:
            if ts is not None:
                content.append({"type": "text", "text": f"[{fmt_ts(ts)}]"})
            uri = f"data:{mime};base64," + base64.b64encode(data).decode()
            content.append({"type": "image_url", "image_url": {"url": uri}})
        content.append({"type": "text", "text": prompt})
        r = self._http.post("/chat/completions", json={
            "model": self.vl_model,
            "messages": [{"role": "user", "content": content}],
            "temperature": 0.1,
        })
        r.raise_for_status()
        return r.json()["choices"][0]["message"]["content"] or ""


# ---------- 文档 ----------
DOC_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".csv"}


def read_document(src: Path) -> str:
    suffix = src.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(src).pages).strip()
    if suffix == ".docx":
        from docx import Document
        doc = Document(src)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            parts += [" | ".join(c.text.strip() for c in row.cells) for row in t.rows]
        return "\n".join(parts)
    raw = src.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


# ---------- ffmpeg 工具 ----------
def media_kind(filename: str, mime: str | None) -> str:
    # 扩展名优先：浏览器常把 .md/.csv/.docx 报成 application/octet-stream
    if Path(filename).suffix.lower() in DOC_SUFFIXES:
        return "document"
    mime = mime or mimetypes.guess_type(filename)[0] or ""
    for k in ("image", "audio", "video"):
        if mime.startswith(k + "/"):
            return k
    return "other"


def _run(cmd: list[str]) -> str:
    p = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    if p.returncode != 0:
        raise RuntimeError(f"{cmd[0]} 失败: {p.stderr[-500:]}")
    return p.stdout


def probe(src: Path) -> dict:
    return json.loads(_run(["ffprobe", "-v", "error", "-print_format", "json",
                            "-show_format", "-show_streams", str(src)]))


def duration(src: Path) -> float:
    return float(probe(src)["format"].get("duration") or 0)


def has_audio(src: Path) -> bool:
    return any(s["codec_type"] == "audio" for s in probe(src)["streams"])


def split_audio(src: Path, tmp: Path) -> list[Path]:
    out = tmp / "seg"
    out.mkdir(exist_ok=True)
    _run(["ffmpeg", "-y", "-v", "error", "-i", str(src), "-vn", "-ac", "1", "-ar", "16000",
          "-b:a", "32k", "-f", "segment", "-segment_time", str(SEGMENT_SECONDS),
          str(out / "%03d.mp3")])
    return sorted(out.glob("*.mp3"))


def extract_frames(src: Path, tmp: Path) -> list[tuple[Path, float]]:
    """均匀抽最多 MAX_FRAMES 帧，至少间隔 5 秒。"""
    dur = duration(src)
    if dur <= 0:
        return []
    step = max(dur / MAX_FRAMES, 5.0)
    out = tmp / "frames"
    out.mkdir(exist_ok=True)
    frames = []
    t = step / 2
    while t < dur and len(frames) < MAX_FRAMES:
        f = out / f"{len(frames):02d}.jpg"
        _run(["ffmpeg", "-y", "-v", "error", "-ss", f"{t:.2f}", "-i", str(src),
              "-frames:v", "1", "-vf", "scale=768:-2", str(f)])
        if f.exists():
            frames.append((f, t))
        t += step
    return frames


def fmt_ts(sec: float) -> str:
    sec = int(sec)
    return f"{sec // 60:02d}:{sec % 60:02d}"
