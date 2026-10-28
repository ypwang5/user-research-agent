"""媒体处理测试：真实 ffmpeg + 模拟的 DashScope 接口。需要 ffmpeg（在容器里跑）。"""

import json
import shutil
import subprocess
from pathlib import Path

import httpx
import pytest

from memory_service.media import MediaProcessor, extract_frames, media_kind, split_audio

pytestmark = pytest.mark.skipif(shutil.which("ffmpeg") is None, reason="需要 ffmpeg")


def ff(*args):
    subprocess.run(["ffmpeg", "-y", "-v", "error", *args], check=True)


class FakeDashScope:
    """记录请求，ASR 返回第 N 段，VL 返回收到的图片数。"""

    def __init__(self, fail_asr=False):
        self.asr_calls = 0
        self.vl_images = []
        self.fail_asr = fail_asr

    def __call__(self, req: httpx.Request) -> httpx.Response:
        body = json.loads(req.content)
        content = body["messages"][0]["content"]
        if body["model"] == "asr":
            if self.fail_asr:
                return httpx.Response(500, json={"error": "boom"})
            self.asr_calls += 1
            assert content[0]["input_audio"]["data"].startswith("data:audio/mpeg;base64,")
            text = f"第{self.asr_calls}段语音"
        else:
            n = sum(c["type"] == "image_url" for c in content)
            self.vl_images.append(n)
            text = f"看到{n}张图"
        return httpx.Response(200, json={"choices": [{"message": {"content": text}}]})


def proc(fake):
    http = httpx.Client(base_url="http://fake", transport=httpx.MockTransport(fake))
    return MediaProcessor("k", "http://fake", "asr", "vl", http=http, workers=1)


def test_media_kind():
    assert media_kind("a.mp3", None) == "audio"
    assert media_kind("a.MP4", None) == "video"
    assert media_kind("x", "image/png") == "image"
    assert media_kind("a.pdf", None) == "document"
    assert media_kind("a.zip", None) == "other"


def test_split_long_audio(tmp_path):
    src = tmp_path / "long.wav"
    ff("-f", "lavfi", "-i", "sine=frequency=440:duration=600", str(src))  # 10 分钟
    segs = split_audio(src, tmp_path)
    assert len(segs) == 3  # 240s 一段 → 3 段
    assert all(s.stat().st_size < 10 * 1024 * 1024 for s in segs)


def test_frames_capped(tmp_path):
    src = tmp_path / "v.mp4"
    ff("-f", "lavfi", "-i", "testsrc=duration=120:size=320x240:rate=5", str(src))
    frames = extract_frames(src, tmp_path)
    assert 1 <= len(frames) <= 8
    assert [round(t) for _, t in frames] == sorted(round(t) for _, t in frames)


def test_audio_transcribed_in_order(tmp_path):
    src = tmp_path / "a.wav"
    ff("-f", "lavfi", "-i", "sine=duration=500", str(src))
    fake = FakeDashScope()
    r = proc(fake).extract("a.wav", src.read_bytes(), "audio/wav")
    assert r["kind"] == "audio"
    assert r["text"].splitlines() == ["[00:00] 第1段语音", "[04:00] 第2段语音", "[08:00] 第3段语音"]


def test_video_with_audio(tmp_path):
    src = tmp_path / "v.mp4"
    ff("-f", "lavfi", "-i", "testsrc=duration=30:size=320x240:rate=5",
       "-f", "lavfi", "-i", "sine=duration=30", "-shortest", str(src))
    fake = FakeDashScope()
    r = proc(fake).extract("v.mp4", src.read_bytes(), "video/mp4")
    assert "【音轨转写】" in r["text"] and "第1段语音" in r["text"]
    assert "【画面描述】" in r["text"]
    assert fake.vl_images and 1 <= fake.vl_images[0] <= 8


def test_silent_video_skips_asr(tmp_path):
    src = tmp_path / "v.mp4"
    ff("-f", "lavfi", "-i", "testsrc=duration=10:size=320x240:rate=5", str(src))
    fake = FakeDashScope()
    r = proc(fake).extract("v.mp4", src.read_bytes(), "video/mp4")
    assert fake.asr_calls == 0
    assert "【画面描述】" in r["text"]


def test_image(tmp_path):
    src = tmp_path / "i.png"
    ff("-f", "lavfi", "-i", "testsrc=size=64x64", "-frames:v", "1", str(src))
    fake = FakeDashScope()
    r = proc(fake).extract("i.png", src.read_bytes(), "image/png")
    assert r["text"] == "看到1张图"


def test_failure_isolated(tmp_path):
    src = tmp_path / "a.wav"
    ff("-f", "lavfi", "-i", "sine=duration=5", str(src))
    r = proc(FakeDashScope(fail_asr=True)).extract("a.wav", src.read_bytes(), "audio/wav")
    assert "error" in r and "text" not in r
    r = proc(FakeDashScope()).extract("x.zip", b"PK", "application/zip")
    assert r["kind"] == "other" and "error" in r


def test_documents(tmp_path):
    from docx import Document
    d = Document()
    d.add_paragraph("受访者说：每月对账要两天")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "平台", "淘宝"
    p = tmp_path / "n.docx"
    d.save(p)
    r = proc(FakeDashScope()).extract("n.docx", p.read_bytes(), "application/octet-stream")
    assert r["kind"] == "document"
    assert "每月对账要两天" in r["text"] and "平台 | 淘宝" in r["text"]

    r = proc(FakeDashScope()).extract("n.txt", "中文笔记".encode("gb18030"), "text/plain")
    assert r["text"] == "中文笔记"
    r = proc(FakeDashScope()).extract("n.md", "# 标题".encode("utf-8"), None)
    assert r["kind"] == "document" and r["text"] == "# 标题"


def test_corrupt_media_reports_error():
    r = proc(FakeDashScope()).extract("bad.mp3", b"not audio", "audio/mpeg")
    assert "error" in r
